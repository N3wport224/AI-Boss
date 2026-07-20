"""Shared file-ingestion infrastructure.

Uploads land in artifacts/ (gitignored — this is runtime data, not project
code, same as orchestrator.db), get SHA-256 hashed for dedupe tracking in the
state store, and are converted into structured data other modules can consume.
"""
import csv
import hashlib
import io
import json
import re
import time
from pathlib import Path
from typing import Optional
from urllib.parse import unquote, urlparse

import httpx
from pypdf import PdfReader

ARTIFACTS_DIR = Path(__file__).resolve().parent.parent / "artifacts"

# Applied to every upload and every file the folder watcher picks up. This is
# meant for the CSV/PDF/JSON files this dashboard is actually for, not
# arbitrary large uploads — both call sites reject before the whole file ever
# lands in memory (a chunked read for uploads, a stat() check for the watcher).
MAX_UPLOAD_BYTES = 20 * 1024 * 1024


class UploadTooLargeError(ValueError):
    def __init__(self, size_bytes: int, limit_bytes: int = MAX_UPLOAD_BYTES):
        self.size_bytes = size_bytes
        self.limit_bytes = limit_bytes
        super().__init__(f"File is {size_bytes} bytes, exceeding the {limit_bytes // (1024 * 1024)} MB limit.")


def ensure_artifacts_dir() -> Path:
    ARTIFACTS_DIR.mkdir(exist_ok=True)
    return ARTIFACTS_DIR


def hash_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


async def read_upload_with_limit(file, max_bytes: Optional[int] = None) -> bytes:
    """Read a FastAPI UploadFile in chunks, aborting as soon as the running
    total exceeds `max_bytes` — so an oversized file is never fully
    buffered into memory just to be rejected.

    `max_bytes` defaults to the *current* value of `MAX_UPLOAD_BYTES` (looked
    up at call time, not bound as a function default) so tests can monkeypatch
    the module-level constant and have callers that don't pass one pick it up.
    """
    if max_bytes is None:
        max_bytes = MAX_UPLOAD_BYTES
    chunks = []
    total = 0
    while True:
        chunk = await file.read(1024 * 1024)
        if not chunk:
            break
        total += len(chunk)
        if total > max_bytes:
            raise UploadTooLargeError(total, max_bytes)
        chunks.append(chunk)
    return b"".join(chunks)


def filename_from_url(url: str) -> str:
    """Best-effort filename for a fetched URL — the last path segment,
    percent-decoded, or a generic fallback if the URL has none (e.g. it
    ends in '/')."""
    name = unquote(Path(urlparse(url).path).name)
    return name or "downloaded_file"


def fetch_url_bytes(url: str, timeout: float = 15.0, max_bytes: Optional[int] = None) -> bytes:
    """GET a URL and return its body, aborting as soon as the running total
    exceeds `max_bytes` — the URL-ingestion counterpart to
    `read_upload_with_limit`'s chunked cap for a direct file upload. Same
    trust model as the `http_request` automation module: the URL is
    whatever the user explicitly supplies, no credentials baked in, so this
    is a plain HTTP fetch rather than a SaaS integration."""
    if max_bytes is None:
        max_bytes = MAX_UPLOAD_BYTES
    total = 0
    chunks = []
    with httpx.Client(timeout=timeout, follow_redirects=True) as client:
        with client.stream("GET", url) as response:
            response.raise_for_status()
            for chunk in response.iter_bytes(1024 * 1024):
                total += len(chunk)
                if total > max_bytes:
                    raise UploadTooLargeError(total, max_bytes)
                chunks.append(chunk)
    return b"".join(chunks)


def save_artifact(filename: str, data: bytes, suffix: str = "") -> Path:
    """Save raw uploaded bytes under a timestamp-prefixed name so repeated
    uploads of same-named files never collide."""
    ensure_artifacts_dir()
    safe_name = f"{int(time.time() * 1000)}_{Path(filename).name}{suffix}"
    path = ARTIFACTS_DIR / safe_name
    path.write_bytes(data)
    return path


def csv_bytes_to_records(data: bytes) -> list[dict]:
    text = data.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    return [dict(row) for row in reader]


def cleanse_records(records: list[dict]) -> tuple[list[dict], dict]:
    """Best-effort cleanup applied to every ingested CSV: trims whitespace from
    keys/values, turns empty-string cells into None, drops fully-blank rows,
    and drops exact-duplicate rows. Deliberately does not guess at column
    types or rewrite real values beyond trimming — no schema inference."""
    stats = {
        "rows_before": len(records),
        "blank_rows_removed": 0,
        "duplicate_rows_removed": 0,
        "cells_trimmed": 0,
        "empty_cells_nulled": 0,
    }
    cleaned = []
    seen = set()

    for row in records:
        new_row = {}
        all_blank = True
        for key, value in row.items():
            clean_key = key.strip() if isinstance(key, str) else key
            if isinstance(value, str):
                trimmed = value.strip()
                if trimmed != value:
                    stats["cells_trimmed"] += 1
                if trimmed == "":
                    new_row[clean_key] = None
                    stats["empty_cells_nulled"] += 1
                else:
                    new_row[clean_key] = trimmed
                    all_blank = False
            else:
                new_row[clean_key] = value
                if value is not None:
                    all_blank = False

        if all_blank:
            stats["blank_rows_removed"] += 1
            continue

        dedupe_key = json.dumps(new_row, sort_keys=True, default=str)
        if dedupe_key in seen:
            stats["duplicate_rows_removed"] += 1
            continue
        seen.add(dedupe_key)
        cleaned.append(new_row)

    stats["rows_after"] = len(cleaned)
    return cleaned, stats


def pdf_bytes_to_text(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def xlsx_bytes_to_records(data: bytes) -> list[dict]:
    """Reads the first sheet of an uploaded .xlsx workbook, treating row 1 as
    headers — the spreadsheet counterpart to csv_bytes_to_records. openpyxl
    is imported lazily (same reason as the XLSX run-history export in
    webapp/main.py) so serving the dashboard never pays for the dependency."""
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rows = workbook.worksheets[0].iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        return []

    headers = [str(cell) if cell is not None else f"column_{i + 1}" for i, cell in enumerate(header_row)]
    records = []
    for row in rows:
        if all(cell is None for cell in row):
            continue
        records.append({headers[i]: (row[i] if i < len(row) else None) for i in range(len(headers))})
    return records


def ingest_csv_bytes(filename: str, data: bytes, store) -> dict:
    """Shared by the upload endpoint and the filesystem watcher: hash-dedupe,
    parse, save the artifact + its JSON conversion, record the hash."""
    file_hash = hash_bytes(data)
    existing = store.find_ingested_file(file_hash)
    if existing:
        return {"duplicate": True, "original_filename": existing["filename"], "ingested_at": existing["ingested_at"]}

    records = csv_bytes_to_records(data)
    records, cleaning_stats = cleanse_records(records)
    saved_path = save_artifact(filename, data)
    saved_path.with_suffix(".json").write_text(json.dumps(records, indent=2))
    store.record_ingested_file(file_hash, filename, "csv")

    return {
        "duplicate": False,
        "filename": filename,
        "row_count": len(records),
        "columns": list(records[0].keys()) if records else [],
        "preview": records[:50],
        "truncated": len(records) > 50,
        "cleaning": cleaning_stats,
    }


def ingest_xlsx_bytes(filename: str, data: bytes, store) -> dict:
    """Shared by the upload endpoint: hash-dedupe, parse the first sheet,
    save the artifact + its JSON conversion, record the hash — the
    spreadsheet counterpart to ingest_csv_bytes, reusing the exact same
    cleanse_records() pass (trim/blank-row/duplicate-row cleanup, no
    schema inference) so an .xlsx upload gets identical treatment to a CSV
    one once it's parsed into records."""
    file_hash = hash_bytes(data)
    existing = store.find_ingested_file(file_hash)
    if existing:
        return {"duplicate": True, "original_filename": existing["filename"], "ingested_at": existing["ingested_at"]}

    records = xlsx_bytes_to_records(data)
    records, cleaning_stats = cleanse_records(records)
    saved_path = save_artifact(filename, data)
    saved_path.with_suffix(".json").write_text(json.dumps(records, indent=2, default=str))
    store.record_ingested_file(file_hash, filename, "xlsx")

    return {
        "duplicate": False,
        "filename": filename,
        "row_count": len(records),
        "columns": list(records[0].keys()) if records else [],
        "preview": records[:50],
        "truncated": len(records) > 50,
        "cleaning": cleaning_stats,
    }


def ingest_pdf_bytes(filename: str, data: bytes, store) -> dict:
    file_hash = hash_bytes(data)
    existing = store.find_ingested_file(file_hash)
    if existing:
        return {"duplicate": True, "original_filename": existing["filename"], "ingested_at": existing["ingested_at"]}

    text = pdf_bytes_to_text(data)
    saved_path = save_artifact(filename, data)
    saved_path.with_suffix(".txt").write_text(text)
    store.record_ingested_file(file_hash, filename, "pdf")

    return {
        "duplicate": False,
        "filename": filename,
        "char_count": len(text),
        "preview": text[:2000],
        "truncated": len(text) > 2000,
    }


def ingest_json_bytes(filename: str, data: bytes, store) -> dict:
    """Shared by the upload endpoint and the filesystem watcher: hash-dedupe,
    parse, save the artifact, record the hash.

    Unlike CSV/PDF, no separate "converted" derivative is written — a JSON
    upload already ships in the structured format ingestion turns CSV/PDF
    into, so the raw uploaded file itself is the reusable artifact (and is
    already picked up by search_artifacts, which looks at .json/.txt files).
    A top-level JSON array is treated as the record list (mirroring a CSV's
    rows); anything else (a single object, a scalar) is wrapped as one row.
    """
    file_hash = hash_bytes(data)
    existing = store.find_ingested_file(file_hash)
    if existing:
        return {"duplicate": True, "original_filename": existing["filename"], "ingested_at": existing["ingested_at"]}

    parsed = json.loads(data.decode("utf-8-sig"))
    records = parsed if isinstance(parsed, list) else [parsed]

    save_artifact(filename, data)
    store.record_ingested_file(file_hash, filename, "json")

    return {
        "duplicate": False,
        "filename": filename,
        "row_count": len(records),
        "columns": list(records[0].keys()) if records and isinstance(records[0], dict) else [],
        "preview": records[:50],
        "truncated": len(records) > 50,
    }


def list_artifacts() -> list[dict]:
    ensure_artifacts_dir()
    files = []
    for path in sorted(ARTIFACTS_DIR.iterdir(), key=lambda p: p.stat().st_mtime, reverse=True):
        if path.is_file():
            stat = path.stat()
            files.append({"name": path.name, "size_bytes": stat.st_size, "modified_at": stat.st_mtime})
    return files


def search_artifacts(query: str, max_results: int = 20) -> list[dict]:
    """Case-insensitive keyword search across every ingested artifact's
    extracted content (.json records for CSVs, .txt text for PDFs) — not the
    raw uploaded bytes. Returns one entry per matching file with a short
    snippet showing where the match was found."""
    ensure_artifacts_dir()
    query_lower = query.lower().strip()
    if not query_lower:
        return []

    results = []
    for path in sorted(ARTIFACTS_DIR.iterdir(), key=lambda p: p.stat().st_mtime, reverse=True):
        if path.suffix not in (".json", ".txt"):
            continue
        try:
            text = path.read_text()
        except (UnicodeDecodeError, OSError):
            continue

        idx = text.lower().find(query_lower)
        if idx == -1:
            continue

        start = max(0, idx - 60)
        end = min(len(text), idx + len(query_lower) + 60)
        snippet = " ".join(text[start:end].split())
        results.append({"artifact": path.name, "kind": path.suffix.lstrip("."), "snippet": snippet})

        if len(results) >= max_results:
            break

    return results


MAX_ARTIFACT_PREVIEW_CHARS = 20_000
MAX_ARTIFACT_PREVIEW_ROWS = 200

_DATE_LIKE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}([T ]\d{2}:\d{2}(:\d{2})?)?$")


def _infer_cell_type(value) -> str:
    """The most specific type a single cell's value looks like -- CSV/JSON
    sidecar records are always strings-or-None (csv.DictReader) or whatever
    JSON scalar type was ingested, so this normalizes both to the same
    vocabulary: int, float, bool, date, text, or null for a blank cell."""
    if value is None or value == "":
        return "null"
    if isinstance(value, bool):
        return "bool"
    if isinstance(value, int):
        return "int"
    if isinstance(value, float):
        return "float"
    text = str(value).strip()
    if text.lower() in ("true", "false"):
        return "bool"
    if _DATE_LIKE_RE.match(text):
        return "date"
    try:
        int(text)
        return "int"
    except ValueError:
        pass
    try:
        float(text)
        return "float"
    except ValueError:
        pass
    return "text"


def infer_schema(records: list[dict]) -> dict:
    """A lightweight column-by-column schema summary for a table-shaped
    artifact -- column names in first-seen order, one inferred type per
    column (the majority type across its non-null cells, "mixed" when no
    type has a clear majority), and the total row count. Deliberately not
    used by cleanse_records() (which stays type-agnostic on purpose) --
    this is read-only, for display in the artifact viewer, never used to
    coerce or rewrite ingested values."""
    columns: list[str] = []
    seen_columns = set()
    for row in records:
        for key in row.keys():
            if key not in seen_columns:
                seen_columns.add(key)
                columns.append(key)

    column_types = {}
    for col in columns:
        counts: dict[str, int] = {}
        for row in records:
            cell_type = _infer_cell_type(row.get(col))
            if cell_type == "null":
                continue
            counts[cell_type] = counts.get(cell_type, 0) + 1
        if not counts:
            column_types[col] = "null"
            continue

        # int is a subset of float for column-typing purposes -- a column of
        # "4" and "3.5" is a float column, not a mixed one.
        numeric_total = counts.get("int", 0) + counts.get("float", 0)
        total = sum(counts.values())
        if numeric_total == total:
            column_types[col] = "float" if counts.get("float", 0) else "int"
            continue

        best_type, best_count = max(counts.items(), key=lambda kv: kv[1])
        column_types[col] = best_type if best_count / total >= 0.9 else "mixed"

    return {
        "columns": [{"name": col, "type": column_types[col]} for col in columns],
        "row_count": len(records),
    }


def read_artifact_content(filename: str) -> dict:
    """Full (but capped) extracted content for one ingested artifact, for
    the dashboard's artifact viewer -- not just a search-result snippet.
    `.json` sidecars (CSV/XLSX records) pretty-print and cap at
    MAX_ARTIFACT_PREVIEW_ROWS records; `.txt` sidecars (PDF text) and raw
    `.csv` uploads cap at MAX_ARTIFACT_PREVIEW_CHARS/rows respectively.
    A binary original (`.pdf`, `.xlsx`) isn't directly viewable as text —
    its `.txt`/`.json` sidecar (ingested alongside it, listed as its own
    artifact) is where the real content lives."""
    path = ARTIFACTS_DIR / filename
    if not path.is_file():
        raise FileNotFoundError(filename)

    suffix = path.suffix.lower()
    if suffix == ".json":
        try:
            data = json.loads(path.read_text())
        except (json.JSONDecodeError, UnicodeDecodeError):
            return {"kind": "text", "content": path.read_text(errors="replace")[:MAX_ARTIFACT_PREVIEW_CHARS], "truncated": True}
        if isinstance(data, list):
            return {
                "kind": "table",
                "content": data[:MAX_ARTIFACT_PREVIEW_ROWS],
                "truncated": len(data) > MAX_ARTIFACT_PREVIEW_ROWS,
                "schema": infer_schema(data) if data and isinstance(data[0], dict) else None,
            }
        return {"kind": "json", "content": data, "truncated": False}

    if suffix == ".txt":
        text = path.read_text(errors="replace")
        return {"kind": "text", "content": text[:MAX_ARTIFACT_PREVIEW_CHARS], "truncated": len(text) > MAX_ARTIFACT_PREVIEW_CHARS}

    if suffix == ".csv":
        records = csv_bytes_to_records(path.read_bytes())
        return {
            "kind": "table",
            "content": records[:MAX_ARTIFACT_PREVIEW_ROWS],
            "truncated": len(records) > MAX_ARTIFACT_PREVIEW_ROWS,
            "schema": infer_schema(records) if records else None,
        }

    return {
        "kind": "unsupported",
        "content": None,
        "truncated": False,
        "message": f"'{suffix}' isn't directly viewable — look for this file's .json or .txt companion instead.",
    }


def purge_old_artifacts(older_than_hours: float) -> list[str]:
    """Delete files under artifacts/ older than the given age.

    Scoped strictly to ARTIFACTS_DIR — never touches orchestrator.db,
    pipelines/, or any source file, no matter what `older_than_hours` is.
    """
    if not ARTIFACTS_DIR.exists():
        return []
    cutoff = time.time() - older_than_hours * 3600
    removed = []
    for path in ARTIFACTS_DIR.iterdir():
        if path.is_file() and path.stat().st_mtime < cutoff:
            path.unlink()
            removed.append(path.name)
    return removed


def delete_artifact(filename: str) -> bool:
    """Delete one specific artifact by name, regardless of age -- the
    finer-grained counterpart to purge_old_artifacts()'s age-based sweep,
    for a checkbox multi-select in the UI. Returns False (a no-op) if the
    file doesn't exist rather than raising."""
    path = ARTIFACTS_DIR / filename
    if not path.is_file():
        return False
    path.unlink()
    return True


class ArtifactRenameError(ValueError):
    """Raised when an artifact rename can't be carried out as requested."""


def rename_artifact(old_name: str, new_name: str) -> None:
    """Rename an ingested artifact's underlying file. Distinct from
    renaming a *tag* (which relabels a tag across every file that carries
    it) -- this moves one specific file to a new name. Only the filename
    itself moves; the caller is responsible for moving the tag-store row
    (and any favorites/recently-viewed keys) to the new name, since those
    live in the state store and localStorage respectively, not on disk."""
    old_path = ARTIFACTS_DIR / old_name
    if not old_path.is_file():
        raise ArtifactRenameError(f"No artifact named '{old_name}'.")
    new_path = ARTIFACTS_DIR / Path(new_name).name
    if new_path == old_path:
        return
    if new_path.exists():
        raise ArtifactRenameError(f"An artifact named '{new_path.name}' already exists.")
    old_path.rename(new_path)
