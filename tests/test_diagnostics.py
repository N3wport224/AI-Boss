import json
import shutil
from datetime import datetime, timedelta

import pytest
from fastapi.testclient import TestClient

from webapp import pipelines as pipeline_store
from webapp.main import app

client = TestClient(app)


@pytest.fixture(autouse=True)
def clean_pipelines_dir():
    if pipeline_store.PIPELINES_DIR.exists():
        shutil.rmtree(pipeline_store.PIPELINES_DIR)
    yield
    if pipeline_store.PIPELINES_DIR.exists():
        shutil.rmtree(pipeline_store.PIPELINES_DIR)


def test_health_check_reports_ready_for_the_bundled_modules():
    res = client.get("/api/health")
    assert res.status_code == 200

    body = res.json()
    assert body["status"] == "ready"
    names = {check["name"] for check in body["checks"]}
    assert names == {"state_store", "manifests", "entrypoints", "environment"}
    assert all(check["ok"] for check in body["checks"])


def test_performance_reports_a_live_process_snapshot():
    res = client.get("/api/performance")
    assert res.status_code == 200

    body = res.json()
    assert body["pid"] > 0
    assert body["memory_rss_mb"] > 0
    assert body["thread_count"] >= 1
    assert body["active_run_threads"] >= 0
    assert body["uptime_seconds"] >= 0


def test_memory_endpoints_list_delete_and_redact():
    from webapp.main import store

    store.set_memory("last_risk_level", "high")
    store.set_memory("agent_secret_token", "sk-should-be-hidden")

    listed = client.get("/api/memory").json()
    entries = {entry["key"]: entry["value"] for entry in listed}
    assert entries["last_risk_level"] == "high"
    assert entries["agent_secret_token"] == "***REDACTED***"

    res = client.delete("/api/memory/last_risk_level")
    assert res.status_code == 200
    assert res.json() == {"deleted": "last_risk_level"}
    remaining_keys = {entry["key"] for entry in client.get("/api/memory").json()}
    assert "last_risk_level" not in remaining_keys

    clear_res = client.delete("/api/memory")
    assert clear_res.status_code == 200
    assert client.get("/api/memory").json() == []


def test_rename_memory_key_moves_the_value_to_the_new_key():
    from webapp.main import store

    store.set_memory("batch41_old_key", {"nested": "value"})

    res = client.post("/api/memory/batch41_old_key/rename", json={"new_key": "batch41_new_key"})
    assert res.status_code == 200
    assert res.json() == {"old_key": "batch41_old_key", "new_key": "batch41_new_key"}

    entries = {entry["key"]: entry["value"] for entry in client.get("/api/memory").json()}
    assert "batch41_old_key" not in entries
    assert entries["batch41_new_key"] == {"nested": "value"}


def test_rename_memory_key_404s_for_an_unknown_key():
    res = client.post("/api/memory/does_not_exist_key/rename", json={"new_key": "whatever"})
    assert res.status_code == 404


def test_rename_memory_key_409s_if_the_new_key_already_exists():
    from webapp.main import store

    store.set_memory("batch41_existing_a", "a")
    store.set_memory("batch41_existing_b", "b")

    res = client.post("/api/memory/batch41_existing_a/rename", json={"new_key": "batch41_existing_b"})
    assert res.status_code == 409


def test_rename_memory_key_rejects_a_blank_new_key():
    from webapp.main import store

    store.set_memory("batch41_blank_source", "x")
    res = client.post("/api/memory/batch41_blank_source/rename", json={"new_key": "   "})
    assert res.status_code == 400


def test_duplicate_memory_key_copies_the_value_and_keeps_the_original():
    from webapp.main import store

    store.set_memory("batch45_dup_source", {"nested": "value"})

    res = client.post("/api/memory/batch45_dup_source/duplicate", json={"new_key": "batch45_dup_target"})
    assert res.status_code == 200
    assert res.json() == {"key": "batch45_dup_source", "new_key": "batch45_dup_target"}

    entries = {entry["key"]: entry["value"] for entry in client.get("/api/memory").json()}
    assert entries["batch45_dup_source"] == {"nested": "value"}
    assert entries["batch45_dup_target"] == {"nested": "value"}


def test_duplicate_memory_key_404s_for_an_unknown_source_key():
    res = client.post("/api/memory/does_not_exist_dup_source/duplicate", json={"new_key": "whatever"})
    assert res.status_code == 404


def test_duplicate_memory_key_409s_if_the_new_key_already_exists():
    from webapp.main import store

    store.set_memory("batch45_dup_existing_a", "a")
    store.set_memory("batch45_dup_existing_b", "b")

    res = client.post("/api/memory/batch45_dup_existing_a/duplicate", json={"new_key": "batch45_dup_existing_b"})
    assert res.status_code == 409


def test_duplicate_memory_key_rejects_a_blank_new_key():
    from webapp.main import store

    store.set_memory("batch45_dup_blank_source", "x")
    res = client.post("/api/memory/batch45_dup_blank_source/duplicate", json={"new_key": "   "})
    assert res.status_code == 400


def test_bulk_delete_memory_keys_removes_only_the_selected_ones():
    from webapp.main import store

    store.set_memory("batch45_bulk_del_a", "a")
    store.set_memory("batch45_bulk_del_b", "b")
    store.set_memory("batch45_bulk_del_keep", "keep")

    res = client.post("/api/memory/bulk-delete", json={"keys": ["batch45_bulk_del_a", "batch45_bulk_del_b"]})
    assert res.status_code == 200
    assert set(res.json()["deleted"]) == {"batch45_bulk_del_a", "batch45_bulk_del_b"}

    remaining_keys = {entry["key"] for entry in client.get("/api/memory").json()}
    assert "batch45_bulk_del_a" not in remaining_keys
    assert "batch45_bulk_del_b" not in remaining_keys
    assert "batch45_bulk_del_keep" in remaining_keys


def test_bulk_delete_memory_keys_skips_an_unknown_key():
    from webapp.main import store

    store.set_memory("batch45_bulk_del_known", "x")

    res = client.post(
        "/api/memory/bulk-delete", json={"keys": ["batch45_bulk_del_known", "does_not_exist_bulk_del"]}
    )
    assert res.status_code == 200
    assert res.json() == {"deleted": ["batch45_bulk_del_known"]}


def test_bulk_delete_memory_keys_is_a_no_op_on_an_empty_selection():
    res = client.post("/api/memory/bulk-delete", json={"keys": []})
    assert res.status_code == 200
    assert res.json() == {"deleted": []}


def test_bulk_export_memory_keys_csv_contains_only_the_selected_rows_and_redacts_secrets():
    from webapp.main import store

    store.set_memory("batch45_bulk_csv_a", "alpha")
    store.set_memory("batch45_bulk_csv_secret_token", "sk-should-be-hidden")
    store.set_memory("batch45_bulk_csv_excluded", "not selected")

    res = client.post(
        "/api/memory/bulk-export-csv",
        json={"keys": ["batch45_bulk_csv_a", "batch45_bulk_csv_secret_token"]},
    )
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("text/csv")
    assert "attachment; filename=memory_selected.csv" in res.headers["content-disposition"]

    lines = res.text.strip().splitlines()
    assert lines[0] == "key,value,updated_at"
    marker_row = next(line for line in lines[1:] if line.startswith("batch45_bulk_csv_a"))
    assert "alpha" in marker_row
    secret_row = next(line for line in lines[1:] if line.startswith("batch45_bulk_csv_secret_token"))
    assert "sk-should-be-hidden" not in secret_row
    assert "REDACTED" in secret_row
    assert not any("batch45_bulk_csv_excluded" in line for line in lines[1:])


def test_bulk_export_memory_keys_csv_is_just_a_header_on_an_empty_selection():
    res = client.post("/api/memory/bulk-export-csv", json={"keys": []})
    assert res.status_code == 200
    lines = res.text.strip().splitlines()
    assert lines == ["key,value,updated_at"]


def test_memory_csv_export_contains_header_and_entries_and_redacts_secrets():
    from webapp.main import store

    store.clear_memory()
    store.set_memory("batch18_marker", "plain value")
    store.set_memory("batch18_secret_token", "sk-should-be-hidden")

    res = client.get("/api/memory.csv")
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("text/csv")
    assert "attachment; filename=memory.csv" in res.headers["content-disposition"]

    lines = res.text.strip().splitlines()
    assert lines[0] == "key,value,updated_at"
    marker_row = next(line for line in lines[1:] if line.startswith("batch18_marker"))
    assert "plain value" in marker_row
    secret_row = next(line for line in lines[1:] if line.startswith("batch18_secret_token"))
    assert "sk-should-be-hidden" not in secret_row
    assert "REDACTED" in secret_row

    store.clear_memory()


def test_memory_csv_export_is_empty_but_valid_with_no_entries():
    from webapp.main import store

    store.clear_memory()
    res = client.get("/api/memory.csv")
    assert res.status_code == 200
    lines = res.text.strip().splitlines()
    assert lines == ["key,value,updated_at"]


# ---- Batch 27: export agent memory to JSON ----


def test_memory_json_export_round_trips_through_import_and_redacts_secrets():
    from webapp.main import store

    store.clear_memory()
    store.set_memory("batch27_marker", "plain value")
    store.set_memory("batch27_secret_token", "sk-should-be-hidden")

    res = client.get("/api/memory.json")
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("application/json")
    assert "attachment; filename=memory.json" in res.headers["content-disposition"]

    body = res.json()
    assert body["batch27_marker"] == "plain value"
    assert body["batch27_secret_token"] != "sk-should-be-hidden"
    assert "REDACTED" in body["batch27_secret_token"]

    store.clear_memory()
    import_res = client.post(
        "/api/memory/import",
        files={"file": ("memory.json", json.dumps(body).encode(), "application/json")},
    )
    assert import_res.status_code == 200
    assert set(import_res.json()["imported"]) == {"batch27_marker", "batch27_secret_token"}

    store.clear_memory()


def test_memory_json_export_is_an_empty_object_with_no_entries():
    from webapp.main import store

    store.clear_memory()
    res = client.get("/api/memory.json")
    assert res.status_code == 200
    assert res.json() == {}


# ---- Batch 20: search across agent memory ----

def test_import_memory_upserts_every_key_from_the_json_file():
    from webapp.main import store

    store.clear_memory()
    store.set_memory("existing_key", "old value")

    payload = json.dumps({"existing_key": "new value", "brand_new_key": {"nested": True}}).encode()
    res = client.post(
        "/api/memory/import",
        files={"file": ("memory_backup.json", payload, "application/json")},
    )
    assert res.status_code == 200
    assert set(res.json()["imported"]) == {"existing_key", "brand_new_key"}

    listed = {entry["key"]: entry["value"] for entry in client.get("/api/memory").json()}
    assert listed["existing_key"] == "new value"
    assert listed["brand_new_key"] == {"nested": True}

    store.clear_memory()


def test_import_memory_rejects_invalid_json():
    res = client.post(
        "/api/memory/import",
        files={"file": ("bad.json", b"not valid json", "application/json")},
    )
    assert res.status_code == 400


def test_import_memory_rejects_a_non_object_top_level_value():
    res = client.post(
        "/api/memory/import",
        files={"file": ("list.json", b'["not", "an", "object"]', "application/json")},
    )
    assert res.status_code == 400


def test_search_memory_finds_a_keyword_in_the_key():
    from webapp.main import store

    store.clear_memory()
    store.set_memory("batch20_marker_key", "some value")

    res = client.get("/api/memory/search", params={"q": "batch20_marker_key"})
    assert res.status_code == 200
    body = res.json()
    assert body["query"] == "batch20_marker_key"
    assert any(r["key"] == "batch20_marker_key" for r in body["results"])
    store.clear_memory()


def test_search_memory_finds_a_keyword_in_the_value():
    from webapp.main import store

    store.clear_memory()
    store.set_memory("some_key", "a batch20 unique value marker")

    res = client.get("/api/memory/search", params={"q": "unique value marker"})
    assert any(r["key"] == "some_key" for r in res.json()["results"])
    store.clear_memory()


def test_search_memory_redacts_a_secret_shaped_key():
    from webapp.main import store

    store.clear_memory()
    store.set_memory("batch20_secret_token", "sk-should-be-hidden")

    res = client.get("/api/memory/search", params={"q": "batch20_secret_token"})
    match = next(r for r in res.json()["results"] if r["key"] == "batch20_secret_token")
    assert match["value"] == "***REDACTED***"
    store.clear_memory()


def test_search_memory_returns_empty_for_a_blank_query():
    res = client.get("/api/memory/search", params={"q": ""})
    assert res.status_code == 200
    assert res.json()["results"] == []


def test_search_memory_finds_nothing_for_an_unmatched_keyword():
    res = client.get("/api/memory/search", params={"q": "zzz_never_used_memory_key_zzz"})
    assert res.json()["results"] == []


def test_prune_runs_removes_only_finished_runs_older_than_cutoff(tmp_path):
    from datetime import datetime, timedelta, timezone

    from engine.context import StepRecord
    from engine.state_store import StateStore

    isolated_store = StateStore(str(tmp_path / "prune_test.db"))

    # An old, finished run -> should be pruned.
    old_run_id = isolated_store.start_run()
    now = datetime.now(timezone.utc)
    isolated_store.log_step(old_run_id, StepRecord("a", "automation", now, now, True, {}))
    isolated_store.finish_run(old_run_id, "completed")
    old_cutoff = (now - timedelta(hours=1)).isoformat()
    isolated_store._conn.execute("UPDATE runs SET finished_at = ? WHERE id = ?", (old_cutoff, old_run_id))
    isolated_store._conn.commit()

    # A recent, finished run -> should survive.
    recent_run_id = isolated_store.start_run()
    isolated_store.finish_run(recent_run_id, "completed")

    removed = isolated_store.prune_runs(older_than_hours=0.01)
    isolated_store.close()

    assert removed == 1

    reopened = StateStore(str(tmp_path / "prune_test.db"))
    remaining_ids = {run["id"] for run in reopened.recent_runs(limit=10)}
    assert old_run_id not in remaining_ids
    assert recent_run_id in remaining_ids
    assert reopened.steps_for_run(old_run_id) == []  # its steps were removed too
    reopened.close()


def test_purge_runs_endpoint_removes_old_runs():
    res = client.post("/api/pipeline/run", json={"inputs": {}})
    with client.stream("GET", f"/api/stream/{res.json()['stream_id']}") as response:
        for line in response.iter_lines():
            if line.startswith("data: ") and '"run_completed"' in line:
                break

    before = client.get("/api/metrics").json()["total_runs"]
    purge_res = client.post("/api/runs/purge", params={"older_than_hours": 0})
    assert purge_res.status_code == 200
    assert purge_res.json()["removed_count"] >= 1


def _run_full_pipeline_to_completion():
    res = client.post("/api/pipeline/run", json={"inputs": {}})
    with client.stream("GET", f"/api/stream/{res.json()['stream_id']}") as response:
        for line in response.iter_lines():
            if line.startswith("data: ") and '"run_completed"' in line:
                break
    return client.get("/api/runs?limit=1").json()[0]["id"]


def test_delete_runs_removes_only_the_chosen_ids(tmp_path):
    from engine.state_store import StateStore

    isolated_store = StateStore(str(tmp_path / "bulk_delete_test.db"))
    keep_id = isolated_store.start_run()
    isolated_store.finish_run(keep_id, "completed")
    delete_id_a = isolated_store.start_run()
    isolated_store.finish_run(delete_id_a, "completed")
    delete_id_b = isolated_store.start_run()
    isolated_store.finish_run(delete_id_b, "failed")

    removed = isolated_store.delete_runs([delete_id_a, delete_id_b])
    isolated_store.close()

    assert removed == 2

    reopened = StateStore(str(tmp_path / "bulk_delete_test.db"))
    remaining_ids = {run["id"] for run in reopened.recent_runs(limit=10)}
    assert remaining_ids == {keep_id}
    reopened.close()


def test_delete_runs_ignores_unknown_ids(tmp_path):
    from engine.state_store import StateStore

    isolated_store = StateStore(str(tmp_path / "bulk_delete_unknown.db"))
    real_id = isolated_store.start_run()
    isolated_store.finish_run(real_id, "completed")

    removed = isolated_store.delete_runs([real_id, 999999])
    isolated_store.close()

    assert removed == 1  # only the real id counted, the fake one silently ignored


def test_delete_runs_with_empty_list_is_a_noop(tmp_path):
    from engine.state_store import StateStore

    isolated_store = StateStore(str(tmp_path / "bulk_delete_empty.db"))
    assert isolated_store.delete_runs([]) == 0
    isolated_store.close()


def test_bulk_delete_runs_endpoint_removes_selected_runs_only():
    run_id_to_delete = _run_full_pipeline_to_completion()
    run_id_to_keep = _run_full_pipeline_to_completion()

    res = client.post("/api/runs/bulk-delete", json={"run_ids": [run_id_to_delete]})
    assert res.status_code == 200
    assert res.json()["removed_count"] == 1

    remaining_ids = {r["id"] for r in client.get("/api/runs?limit=50").json()}
    assert run_id_to_delete not in remaining_ids
    assert run_id_to_keep in remaining_ids


def test_bulk_delete_runs_endpoint_handles_empty_selection():
    res = client.post("/api/runs/bulk-delete", json={"run_ids": []})
    assert res.status_code == 200
    assert res.json()["removed_count"] == 0


def test_restore_snapshot_is_additive_and_idempotent(tmp_path):
    from engine.state_store import StateStore

    source_store = StateStore(str(tmp_path / "source.db"))
    run_id = source_store.start_run()
    source_store.finish_run(run_id, "completed")
    source_store.set_memory("some_key", {"nested": "value"})
    snapshot = source_store.export_snapshot()
    source_store.close()

    target_store = StateStore(str(tmp_path / "target.db"))
    first_counts = target_store.restore_snapshot(snapshot)
    assert first_counts["runs"] == 1
    assert first_counts["memory"] == 1

    assert target_store.get_memory("some_key") == {"nested": "value"}
    assert {r["id"] for r in target_store.recent_runs(10)} == {run_id}

    # Restoring the exact same snapshot again inserts nothing new.
    second_counts = target_store.restore_snapshot(snapshot)
    assert second_counts == {"runs": 0, "steps": 0, "schedules": 0, "ingested_files": 0, "memory": 0}

    target_store.close()


def test_restore_snapshot_never_overwrites_an_existing_memory_key():
    from webapp.main import store

    store.set_memory("protected_key", "original_value")
    fake_snapshot = {"memory": [{"key": "protected_key", "value": "attempted_overwrite", "updated_at": None}]}

    store.restore_snapshot(fake_snapshot)
    assert store.get_memory("protected_key") == "original_value"


def test_backup_restore_endpoint_fills_in_missing_runs_and_memory():
    from webapp.main import store

    fake_run_id = 987654321
    fake_snapshot = {
        "runs": [{"id": fake_run_id, "started_at": "2020-01-01T00:00:00+00:00", "finished_at": "2020-01-01T00:00:01+00:00", "status": "completed"}],
        "steps": [],
        "schedules": [],
        "ingested_files": [],
        "memory": [{"key": "restored_test_key", "value": "restored_value", "updated_at": None}],
        "pipelines": [],
    }
    snapshot_bytes = json.dumps(fake_snapshot).encode()

    res = client.post("/api/backup/restore", files={"file": ("backup.json", snapshot_bytes, "application/json")})
    assert res.status_code == 200
    body = res.json()
    assert body["runs"] == 1
    assert body["memory"] == 1
    assert body["pipelines"] == 0

    assert store.get_memory("restored_test_key") == "restored_value"
    run_ids = {r["id"] for r in client.get("/api/runs?limit=1000").json()}
    assert fake_run_id in run_ids

    # Restoring the same snapshot again is a no-op — nothing new inserted.
    res2 = client.post("/api/backup/restore", files={"file": ("backup.json", snapshot_bytes, "application/json")})
    assert res2.json()["runs"] == 0
    assert res2.json()["memory"] == 0


def test_backup_restore_rejects_malformed_json():
    res = client.post("/api/backup/restore", files={"file": ("bad.json", b"{not valid", "application/json")})
    assert res.status_code == 400


def test_state_store_export_snapshot_reflects_runs_and_steps(tmp_path):
    from datetime import datetime, timezone

    from engine.context import StepRecord
    from engine.state_store import StateStore

    isolated_store = StateStore(str(tmp_path / "isolated.db"))
    run_id = isolated_store.start_run()
    now = datetime.now(timezone.utc)
    isolated_store.log_step(run_id, StepRecord("fetch_raw_metrics", "automation", now, now, True, {"raw_metrics": {"signups": 5}}))
    isolated_store.finish_run(run_id, "completed")

    snapshot = isolated_store.export_snapshot()
    isolated_store.close()

    assert len(snapshot["runs"]) == 1
    assert snapshot["runs"][0]["status"] == "completed"
    assert len(snapshot["steps"]) == 1
    assert snapshot["steps"][0]["output"] == {"raw_metrics": {"signups": 5}}
    assert snapshot["steps"][0]["success"] is True
    assert snapshot["schedules"] == []
    assert snapshot["ingested_files"] == []


def test_backup_export_includes_a_completed_run_and_saved_pipelines():
    res = client.post("/api/pipeline/run", json={"inputs": {}})
    with client.stream("GET", f"/api/stream/{res.json()['stream_id']}") as response:
        for line in response.iter_lines():
            if line.startswith("data: ") and '"run_completed"' in line:
                break

    client.post(
        "/api/pipelines",
        json={"name": "Backup Test", "steps": [{"tier": "automation", "name": "fetch_raw_metrics"}]},
    )

    res = client.get("/api/backup/export")
    assert res.status_code == 200
    body = res.json()

    assert "exported_at" in body
    assert len(body["runs"]) >= 1
    assert len(body["steps"]) >= 1
    # A step's output is parsed back into a real dict, not left as a JSON string.
    assert isinstance(body["steps"][0]["output"], dict)
    assert any(p["slug"] == "backup_test" for p in body["pipelines"])

    from webapp import pipelines as pipeline_store
    import shutil

    shutil.rmtree(pipeline_store.PIPELINES_DIR, ignore_errors=True)


def test_backup_db_download_is_a_valid_sqlite_file(tmp_path):
    res = client.get("/api/backup/db")
    assert res.status_code == 200
    assert res.headers["content-disposition"].endswith('filename="orchestrator.db"')

    downloaded = tmp_path / "downloaded.db"
    downloaded.write_bytes(res.content)

    import sqlite3

    conn = sqlite3.connect(str(downloaded))
    tables = {row[0] for row in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    conn.close()
    assert {"runs", "steps", "schedules"}.issubset(tables)


# ---- Batch 31: automatic periodic backup snapshot to disk ----


def test_auto_backup_status_defaults_to_disabled():
    from webapp.main import _auto_backup

    _auto_backup.configure(enabled=False, interval_hours=24.0, keep_count=7)
    res = client.get("/api/backup/auto/status")
    assert res.status_code == 200
    body = res.json()
    assert body["enabled"] is False
    assert body["interval_hours"] == 24.0
    assert body["keep_count"] == 7


def test_configure_auto_backup_rejects_invalid_values():
    res = client.patch("/api/backup/auto", json={"enabled": True, "interval_hours": 0, "keep_count": 3})
    assert res.status_code == 400
    res2 = client.patch("/api/backup/auto", json={"enabled": True, "interval_hours": 1, "keep_count": 0})
    assert res2.status_code == 400


def test_configure_and_run_auto_backup_writes_a_snapshot_file():
    from webapp.main import BACKUPS_DIR, _auto_backup
    import shutil

    shutil.rmtree(BACKUPS_DIR, ignore_errors=True)
    try:
        res = client.patch("/api/backup/auto", json={"enabled": True, "interval_hours": 12, "keep_count": 4})
        assert res.status_code == 200
        assert res.json() == {
            "enabled": True,
            "interval_hours": 12.0,
            "keep_count": 4,
            "last_backup_at": None,
            "next_backup_at": None,
        }

        run_res = client.post("/api/backup/auto/run-now")
        assert run_res.status_code == 200
        assert run_res.json()["backed_up_at"] is not None

        files = list(BACKUPS_DIR.glob("backup_*.json"))
        assert len(files) == 1

        status_res = client.get("/api/backup/auto/status")
        status_body = status_res.json()
        assert status_body["last_backup_at"] is not None
        assert status_body["next_backup_at"] is not None
    finally:
        _auto_backup.configure(enabled=False, interval_hours=24.0, keep_count=7)
        shutil.rmtree(BACKUPS_DIR, ignore_errors=True)


def test_list_auto_backups_is_empty_when_no_backups_dir_exists():
    from webapp.main import BACKUPS_DIR
    import shutil

    shutil.rmtree(BACKUPS_DIR, ignore_errors=True)
    res = client.get("/api/backup/auto/list")
    assert res.status_code == 200
    assert res.json() == []


def test_list_and_restore_from_an_automatic_backup_snapshot_file():
    from webapp.main import BACKUPS_DIR, _auto_backup
    import shutil

    shutil.rmtree(BACKUPS_DIR, ignore_errors=True)
    try:
        run_res = client.post("/api/backup/auto/run-now")
        assert run_res.status_code == 200

        list_res = client.get("/api/backup/auto/list")
        assert list_res.status_code == 200
        snapshots = list_res.json()
        assert len(snapshots) == 1
        assert snapshots[0]["filename"].startswith("backup_")
        assert snapshots[0]["size_bytes"] > 0
        assert snapshots[0]["modified_at"]

        restore_res = client.post(f"/api/backup/auto/restore/{snapshots[0]['filename']}")
        assert restore_res.status_code == 200
        assert "runs" in restore_res.json()
    finally:
        _auto_backup.configure(enabled=False, interval_hours=24.0, keep_count=7)
        shutil.rmtree(BACKUPS_DIR, ignore_errors=True)


def test_restore_auto_backup_404s_for_an_unknown_or_unsafe_filename():
    res = client.post("/api/backup/auto/restore/does-not-exist.json")
    assert res.status_code == 404

    res = client.post("/api/backup/auto/restore/..%2F..%2Fetc%2Fpasswd")
    assert res.status_code == 404


def test_delete_an_individual_automatic_backup_snapshot():
    from webapp.main import BACKUPS_DIR, _auto_backup
    import shutil

    shutil.rmtree(BACKUPS_DIR, ignore_errors=True)
    try:
        client.post("/api/backup/auto/run-now")
        files = list(BACKUPS_DIR.glob("backup_*.json"))
        assert len(files) == 1
        filename = files[0].name

        res = client.delete(f"/api/backup/auto/snapshot/{filename}")
        assert res.status_code == 200
        assert res.json() == {"deleted": filename}
        assert not (BACKUPS_DIR / filename).exists()

        list_res = client.get("/api/backup/auto/list")
        assert list_res.json() == []
    finally:
        _auto_backup.configure(enabled=False, interval_hours=24.0, keep_count=7)
        shutil.rmtree(BACKUPS_DIR, ignore_errors=True)


def test_delete_auto_backup_snapshot_404s_for_an_unknown_or_unsafe_filename():
    res = client.delete("/api/backup/auto/snapshot/does-not-exist.json")
    assert res.status_code == 404

    res = client.delete("/api/backup/auto/snapshot/..%2F..%2Fetc%2Fpasswd")
    assert res.status_code == 404


def test_download_an_individual_automatic_backup_snapshot_file():
    from webapp.main import BACKUPS_DIR, _auto_backup
    import shutil

    shutil.rmtree(BACKUPS_DIR, ignore_errors=True)
    try:
        client.post("/api/backup/auto/run-now")
        files = list(BACKUPS_DIR.glob("backup_*.json"))
        assert len(files) == 1
        filename = files[0].name

        res = client.get(f"/api/backup/auto/snapshot/{filename}")
        assert res.status_code == 200
        assert res.headers["content-type"].startswith("application/json")
        assert filename in res.headers["content-disposition"]
        assert res.json() == json.loads(files[0].read_text())
    finally:
        _auto_backup.configure(enabled=False, interval_hours=24.0, keep_count=7)
        shutil.rmtree(BACKUPS_DIR, ignore_errors=True)


def test_download_auto_backup_snapshot_404s_for_an_unknown_or_unsafe_filename():
    res = client.get("/api/backup/auto/snapshot/does-not-exist.json")
    assert res.status_code == 404

    res = client.get("/api/backup/auto/snapshot/..%2F..%2Fetc%2Fpasswd")
    assert res.status_code == 404


def test_protect_and_unprotect_an_automatic_backup_snapshot():
    from webapp.main import BACKUPS_DIR, _auto_backup
    import shutil

    shutil.rmtree(BACKUPS_DIR, ignore_errors=True)
    try:
        client.post("/api/backup/auto/run-now")
        filename = next(BACKUPS_DIR.glob("backup_*.json")).name

        listed = client.get("/api/backup/auto/list").json()
        assert listed[0]["protected"] is False

        res = client.put(f"/api/backup/auto/snapshot/{filename}/protect", json={"protected": True})
        assert res.status_code == 200
        assert res.json() == {"filename": filename, "protected": True}

        listed = client.get("/api/backup/auto/list").json()
        assert listed[0]["protected"] is True

        res2 = client.put(f"/api/backup/auto/snapshot/{filename}/protect", json={"protected": False})
        assert res2.json() == {"filename": filename, "protected": False}
        listed = client.get("/api/backup/auto/list").json()
        assert listed[0]["protected"] is False
    finally:
        _auto_backup.configure(enabled=False, interval_hours=24.0, keep_count=7)
        shutil.rmtree(BACKUPS_DIR, ignore_errors=True)


def test_protect_auto_backup_snapshot_404s_for_an_unknown_or_unsafe_filename():
    res = client.put("/api/backup/auto/snapshot/does-not-exist.json/protect", json={"protected": True})
    assert res.status_code == 404

    res = client.put("/api/backup/auto/snapshot/..%2F..%2Fetc%2Fpasswd/protect", json={"protected": True})
    assert res.status_code == 404


def test_deleting_a_protected_snapshot_still_succeeds_and_clears_its_protection_row():
    from webapp.main import BACKUPS_DIR, _auto_backup, store
    import shutil

    shutil.rmtree(BACKUPS_DIR, ignore_errors=True)
    try:
        client.post("/api/backup/auto/run-now")
        filename = next(BACKUPS_DIR.glob("backup_*.json")).name
        client.put(f"/api/backup/auto/snapshot/{filename}/protect", json={"protected": True})
        assert store.is_backup_protected(filename)

        res = client.delete(f"/api/backup/auto/snapshot/{filename}")
        assert res.status_code == 200
        assert not (BACKUPS_DIR / filename).exists()
        assert not store.is_backup_protected(filename)
    finally:
        _auto_backup.configure(enabled=False, interval_hours=24.0, keep_count=7)
        shutil.rmtree(BACKUPS_DIR, ignore_errors=True)


def test_bulk_protect_and_unprotect_selected_automatic_backup_snapshots():
    import time as _time
    from webapp.main import BACKUPS_DIR, _auto_backup
    import shutil

    shutil.rmtree(BACKUPS_DIR, ignore_errors=True)
    try:
        client.post("/api/backup/auto/run-now")
        _time.sleep(1.1)  # filenames are second-resolution timestamps -- force distinct names
        client.post("/api/backup/auto/run-now")
        files = sorted(BACKUPS_DIR.glob("backup_*.json"))
        assert len(files) == 2
        a_name, b_name = files[0].name, files[1].name

        res = client.post(
            "/api/backup/auto/bulk-protect",
            json={"filenames": [a_name, b_name, "does-not-exist.json"], "protected": True},
        )
        assert res.status_code == 200
        body = res.json()
        assert body["protected"] is True
        assert set(body["updated"]) == {a_name, b_name}

        listed = {s["filename"]: s for s in client.get("/api/backup/auto/list").json()}
        assert listed[a_name]["protected"] is True
        assert listed[b_name]["protected"] is True

        res2 = client.post("/api/backup/auto/bulk-protect", json={"filenames": [a_name], "protected": False})
        assert res2.json()["updated"] == [a_name]

        listed2 = {s["filename"]: s for s in client.get("/api/backup/auto/list").json()}
        assert listed2[a_name]["protected"] is False
        assert listed2[b_name]["protected"] is True
    finally:
        _auto_backup.configure(enabled=False, interval_hours=24.0, keep_count=7)
        shutil.rmtree(BACKUPS_DIR, ignore_errors=True)


def test_bulk_protect_auto_backups_is_a_no_op_on_an_empty_selection():
    res = client.post("/api/backup/auto/bulk-protect", json={"filenames": [], "protected": True})
    assert res.status_code == 200
    assert res.json() == {"protected": True, "updated": []}


def test_auto_backup_list_csv_export_has_a_header_and_matches_the_json_list():
    from webapp.main import BACKUPS_DIR, _auto_backup
    import shutil

    shutil.rmtree(BACKUPS_DIR, ignore_errors=True)
    try:
        client.post("/api/backup/auto/run-now")
        json_list = client.get("/api/backup/auto/list").json()
        assert len(json_list) == 1

        csv_res = client.get("/api/backup/auto/list.csv")
        assert csv_res.status_code == 200
        assert csv_res.headers["content-type"].startswith("text/csv")

        lines = csv_res.text.strip().splitlines()
        assert lines[0] == "filename,size_bytes,modified_at,protected"
        assert len(lines) == 2
        assert lines[1].startswith(json_list[0]["filename"])
    finally:
        _auto_backup.configure(enabled=False, interval_hours=24.0, keep_count=7)
        shutil.rmtree(BACKUPS_DIR, ignore_errors=True)


def test_auto_backup_list_csv_export_is_just_a_header_when_empty():
    from webapp.main import BACKUPS_DIR
    import shutil

    shutil.rmtree(BACKUPS_DIR, ignore_errors=True)
    res = client.get("/api/backup/auto/list.csv")
    assert res.status_code == 200
    assert res.text.strip().splitlines() == ["filename,size_bytes,modified_at,protected"]


def test_download_all_auto_backups_bundles_every_snapshot_into_a_zip():
    import io
    import zipfile

    from webapp.main import BACKUPS_DIR, _auto_backup
    import shutil

    shutil.rmtree(BACKUPS_DIR, ignore_errors=True)
    try:
        client.post("/api/backup/auto/run-now")
        import time as _time

        _time.sleep(1.1)  # filenames are second-resolution timestamps -- force a distinct second file
        client.post("/api/backup/auto/run-now")

        json_list = client.get("/api/backup/auto/list").json()
        assert len(json_list) == 2

        res = client.get("/api/backup/auto/download-all")
        assert res.status_code == 200
        assert res.headers["content-type"] == "application/zip"
        assert res.headers["content-disposition"] == "attachment; filename=auto_backup_snapshots.zip"

        zf = zipfile.ZipFile(io.BytesIO(res.content))
        assert set(zf.namelist()) == {s["filename"] for s in json_list}
    finally:
        _auto_backup.configure(enabled=False, interval_hours=24.0, keep_count=7)
        shutil.rmtree(BACKUPS_DIR, ignore_errors=True)


def test_download_all_auto_backups_404s_when_there_are_none():
    from webapp.main import BACKUPS_DIR
    import shutil

    shutil.rmtree(BACKUPS_DIR, ignore_errors=True)
    res = client.get("/api/backup/auto/download-all")
    assert res.status_code == 404


def test_purge_auto_backups_deletes_snapshots_older_than_the_cutoff():
    import re

    from webapp.main import BACKUPS_DIR, _auto_backup
    import shutil

    shutil.rmtree(BACKUPS_DIR, ignore_errors=True)
    try:
        client.post("/api/backup/auto/run-now")
        files = list(BACKUPS_DIR.glob("backup_*.json"))
        assert len(files) == 1

        # Rewrite the snapshot's own filename to look 48 hours old, since
        # purge reads the cutoff from the filename's encoded timestamp.
        match = re.match(r"backup_(\d{8})T(\d{6})Z\.json", files[0].name)
        stale_time = (
            datetime.strptime(match.group(1) + match.group(2), "%Y%m%d%H%M%S") - timedelta(hours=48)
        ).strftime("%Y%m%dT%H%M%SZ")
        files[0].rename(files[0].with_name(f"backup_{stale_time}.json"))

        res = client.post("/api/backup/auto/purge?older_than_hours=24")
        assert res.status_code == 200
        assert res.json() == {"deleted": 1}
        assert list(BACKUPS_DIR.glob("backup_*.json")) == []
    finally:
        _auto_backup.configure(enabled=False, interval_hours=24.0, keep_count=7)
        shutil.rmtree(BACKUPS_DIR, ignore_errors=True)


def test_purge_auto_backups_rejects_a_non_positive_cutoff():
    res = client.post("/api/backup/auto/purge?older_than_hours=0")
    assert res.status_code == 400


def test_compare_auto_backups_reports_a_delta_for_an_added_run():
    import time

    from webapp.main import BACKUPS_DIR, _auto_backup
    import shutil

    shutil.rmtree(BACKUPS_DIR, ignore_errors=True)
    try:
        client.post("/api/backup/auto/run-now")
        files_before = list(BACKUPS_DIR.glob("backup_*.json"))
        assert len(files_before) == 1
        filename_a = files_before[0].name

        client.post("/api/pipeline/run", json={"inputs": {}})  # add a new run in between

        time.sleep(1.1)  # filenames are second-resolution timestamps -- force a distinct second file
        client.post("/api/backup/auto/run-now")
        filename_b = next(
            f.name for f in BACKUPS_DIR.glob("backup_*.json") if f.name != filename_a
        )

        res = client.get(f"/api/backup/auto/compare?a={filename_a}&b={filename_b}")
        assert res.status_code == 200
        body = res.json()
        assert body["a"]["filename"] == filename_a
        assert body["b"]["filename"] == filename_b
        assert body["a"]["exported_at"]
        assert body["b"]["exported_at"]
        assert body["delta"]["runs"] >= 1
        assert body["b"]["counts"]["runs"] - body["a"]["counts"]["runs"] == body["delta"]["runs"]
    finally:
        _auto_backup.configure(enabled=False, interval_hours=24.0, keep_count=7)
        shutil.rmtree(BACKUPS_DIR, ignore_errors=True)


def test_compare_auto_backups_404s_for_an_unknown_or_unsafe_filename():
    res = client.get("/api/backup/auto/compare?a=does-not-exist.json&b=also-missing.json")
    assert res.status_code == 404


def test_a_failing_automatic_backup_tick_records_a_backup_failed_notification():
    from webapp.main import _auto_backup

    broken = RuntimeError("disk is full (simulated)")
    original_build_snapshot = _auto_backup.build_snapshot
    _auto_backup.build_snapshot = lambda: (_ for _ in ()).throw(broken)
    # interval_hours is set absurdly small (rather than a realistic value) so
    # this tick is due regardless of whatever last_backup_at another test in
    # this same process may have already left on the shared singleton.
    _auto_backup.configure(enabled=True, interval_hours=0.0000001, keep_count=7)
    last_backup_at_before = _auto_backup.last_backup_at
    try:
        _auto_backup._tick()
        notifications = client.get("/api/notifications?limit=20").json()
        assert any(
            n["kind"] == "backup_failed" and "disk is full (simulated)" in n["message"] for n in notifications
        )
        assert _auto_backup.last_backup_at == last_backup_at_before, "a failed write must not be recorded as a successful backup"
    finally:
        _auto_backup.build_snapshot = original_build_snapshot
        _auto_backup.configure(enabled=False, interval_hours=24.0, keep_count=7)


def test_metrics_reflect_a_completed_run():
    before = client.get("/api/metrics").json()

    res = client.post("/api/pipeline/run", json={"inputs": {}})
    stream_id = res.json()["stream_id"]
    with client.stream("GET", f"/api/stream/{stream_id}") as response:
        for line in response.iter_lines():
            if line.startswith("data: ") and '"run_completed"' in line:
                break

    after = client.get("/api/metrics").json()
    assert after["total_runs"] == before["total_runs"] + 1
    assert after["completed"] >= 1
    assert after["success_rate"] is not None
    assert after["avg_duration_seconds"] is not None


def test_runs_csv_export_has_a_header_and_rows():
    client.post("/api/pipeline/run", json={"inputs": {}})
    res = client.get("/api/runs.csv")
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("text/csv")

    lines = res.text.strip().splitlines()
    assert lines[0] == "id,started_at,finished_at,status,duration_seconds"
    assert len(lines) >= 2


def test_runs_json_export_matches_the_summary_list_shape():
    client.post("/api/pipeline/run", json={"inputs": {}})
    res = client.get("/api/runs.json")
    assert res.status_code == 200
    assert "application/json" in res.headers["content-type"]
    assert res.headers["content-disposition"] == "attachment; filename=run_history.json"

    runs = res.json()
    assert isinstance(runs, list)
    assert len(runs) >= 1
    assert set(runs[0].keys()) == {"id", "started_at", "finished_at", "status", "note"}


def test_runs_xlsx_export_carries_runs_and_per_step_detail():
    import io

    from openpyxl import load_workbook

    # Run something first so both sheets have at least one data row.
    res = client.post(
        "/api/modules/automation/fetch_raw_metrics/run",
        json={"inputs": {"signups": 5, "churn": 1, "revenue": 10}, "force_refresh": True},
    )
    with client.stream("GET", f"/api/stream/{res.json()['stream_id']}") as response:
        for line in response.iter_lines():
            if line.startswith("data: ") and '"run_completed"' in line:
                break

    res = client.get("/api/runs.xlsx")
    assert res.status_code == 200
    assert "spreadsheetml" in res.headers["content-type"]

    workbook = load_workbook(io.BytesIO(res.content))
    assert workbook.sheetnames == ["Runs", "Steps"]

    runs_rows = list(workbook["Runs"].values)
    assert runs_rows[0] == ("id", "started_at", "finished_at", "status", "duration_seconds")
    assert len(runs_rows) >= 2

    steps_rows = list(workbook["Steps"].values)
    assert steps_rows[0] == ("run_id", "step", "tier", "success", "error", "started_at", "finished_at")
    assert len(steps_rows) >= 2
    # Every step row's run_id points at a run present in the Runs sheet.
    run_ids = {row[0] for row in runs_rows[1:]}
    assert all(row[0] in run_ids for row in steps_rows[1:])


def test_duplicate_pipeline_creates_a_distinct_copy():
    client.post(
        "/api/pipelines",
        json={"name": "Original", "steps": [{"tier": "automation", "name": "fetch_raw_metrics"}]},
    )

    res = client.post("/api/pipelines/original/duplicate")
    assert res.status_code == 200
    duplicated = res.json()["pipeline"]
    assert duplicated["slug"] == "original_copy"
    assert duplicated["name"] == "Original (copy)"
    assert duplicated["steps"] == [
        {"tier": "automation", "name": "fetch_raw_metrics", "inputs": {}, "mappings": {}}
    ]

    slugs = {p["slug"] for p in client.get("/api/pipelines").json()}
    assert slugs == {"original", "original_copy"}


def test_duplicate_unknown_pipeline_404s():
    res = client.post("/api/pipelines/does_not_exist/duplicate")
    assert res.status_code == 404


def test_rename_pipeline_moves_the_slug_and_keeps_the_steps():
    client.post(
        "/api/pipelines",
        json={"name": "Rename Me", "steps": [{"tier": "automation", "name": "fetch_raw_metrics"}]},
    )

    res = client.post("/api/pipelines/rename_me/rename", json={"name": "Renamed Pipeline"})
    assert res.status_code == 200
    renamed = res.json()["pipeline"]
    assert renamed["slug"] == "renamed_pipeline"
    assert renamed["name"] == "Renamed Pipeline"
    assert renamed["steps"] == [
        {"tier": "automation", "name": "fetch_raw_metrics", "inputs": {}, "mappings": {}}
    ]

    slugs = {p["slug"] for p in client.get("/api/pipelines").json()}
    assert "rename_me" not in slugs
    assert "renamed_pipeline" in slugs


def test_rename_pipeline_to_the_same_slug_only_changes_display_name():
    client.post(
        "/api/pipelines",
        json={"name": "Case Test", "steps": [{"tier": "automation", "name": "fetch_raw_metrics"}]},
    )

    res = client.post("/api/pipelines/case_test/rename", json={"name": "CASE TEST"})
    assert res.status_code == 200
    renamed = res.json()["pipeline"]
    assert renamed["slug"] == "case_test"
    assert renamed["name"] == "CASE TEST"


def test_rename_pipeline_migrates_tags_and_repoints_schedules():
    client.post(
        "/api/pipelines",
        json={"name": "Migrate Me", "steps": [{"tier": "automation", "name": "fetch_raw_metrics"}]},
    )
    client.put("/api/pipelines/migrate_me/tags", json={"tags": ["important"]})
    schedule_res = client.post(
        "/api/schedules",
        json={"kind": "pipeline", "name": "migrate_me", "interval_seconds": 3600},
    )
    schedule_id = schedule_res.json()["id"]

    res = client.post("/api/pipelines/migrate_me/rename", json={"name": "Migrated Pipeline"})
    assert res.status_code == 200

    assert client.get("/api/pipelines").json()
    pipelines_by_slug = {p["slug"]: p for p in client.get("/api/pipelines").json()}
    assert pipelines_by_slug["migrated_pipeline"]["tags"] == ["important"]
    assert "migrate_me" not in pipelines_by_slug

    schedules = client.get("/api/schedules").json()
    schedule = next(s for s in schedules if s["id"] == schedule_id)
    assert schedule["name"] == "migrated_pipeline"

    client.delete(f"/api/schedules/{schedule_id}")


def test_rename_pipeline_rejects_a_collision_with_an_existing_pipeline():
    client.post(
        "/api/pipelines",
        json={"name": "Rename Collision A", "steps": [{"tier": "automation", "name": "fetch_raw_metrics"}]},
    )
    client.post(
        "/api/pipelines",
        json={"name": "Rename Collision B", "steps": [{"tier": "automation", "name": "fetch_raw_metrics"}]},
    )

    res = client.post("/api/pipelines/rename_collision_a/rename", json={"name": "Rename Collision B"})
    assert res.status_code == 400

    slugs = {p["slug"] for p in client.get("/api/pipelines").json()}
    assert "rename_collision_a" in slugs
    assert "rename_collision_b" in slugs


def test_rename_unknown_pipeline_404s():
    res = client.post("/api/pipelines/does_not_exist/rename", json={"name": "Whatever"})
    assert res.status_code == 404


def test_bulk_duplicate_pipelines_clones_every_selected_pipeline():
    client.post(
        "/api/pipelines",
        json={"name": "Bulk Dup A", "steps": [{"tier": "automation", "name": "fetch_raw_metrics"}]},
    )
    client.post(
        "/api/pipelines",
        json={"name": "Bulk Dup B", "steps": [{"tier": "automation", "name": "fetch_raw_metrics"}]},
    )

    res = client.post("/api/pipelines/bulk-duplicate", json={"slugs": ["bulk_dup_a", "bulk_dup_b"]})
    assert res.status_code == 200
    duplicated = res.json()["duplicated"]
    assert {d["slug"] for d in duplicated} == {"bulk_dup_a_copy", "bulk_dup_b_copy"}

    slugs = {p["slug"] for p in client.get("/api/pipelines").json()}
    assert {"bulk_dup_a", "bulk_dup_a_copy", "bulk_dup_b", "bulk_dup_b_copy"} <= slugs


def test_bulk_duplicate_pipelines_skips_unknown_slugs_without_failing():
    client.post(
        "/api/pipelines",
        json={"name": "Bulk Dup C", "steps": [{"tier": "automation", "name": "fetch_raw_metrics"}]},
    )

    res = client.post(
        "/api/pipelines/bulk-duplicate",
        json={"slugs": ["bulk_dup_c", "totally_made_up_slug"]},
    )
    assert res.status_code == 200
    duplicated = res.json()["duplicated"]
    assert [d["slug"] for d in duplicated] == ["bulk_dup_c_copy"]


def test_bulk_duplicate_pipelines_is_a_no_op_on_an_empty_selection():
    res = client.post("/api/pipelines/bulk-duplicate", json={"slugs": []})
    assert res.status_code == 200
    assert res.json() == {"duplicated": []}


def test_bulk_tag_pipelines_adds_the_tag_to_every_selected_pipeline():
    client.post(
        "/api/pipelines",
        json={"name": "Bulk Tag A", "steps": [{"tier": "automation", "name": "fetch_raw_metrics"}]},
    )
    client.post(
        "/api/pipelines",
        json={"name": "Bulk Tag B", "steps": [{"tier": "automation", "name": "fetch_raw_metrics"}]},
    )
    client.put("/api/pipelines/bulk_tag_a/tags", json={"tags": ["existing"]})

    res = client.post("/api/pipelines/bulk-tags", json={"slugs": ["bulk_tag_a", "bulk_tag_b"], "tag": "new_tag"})
    assert res.status_code == 200
    body = res.json()
    assert body["tag"] == "new_tag"
    assert set(body["tagged"]) == {"bulk_tag_a", "bulk_tag_b"}

    updated = {p["slug"]: p for p in client.get("/api/pipelines").json()}
    assert set(updated["bulk_tag_a"]["tags"]) == {"existing", "new_tag"}
    assert updated["bulk_tag_b"]["tags"] == ["new_tag"]


def test_bulk_tag_pipelines_skips_unknown_slugs_without_failing():
    client.post(
        "/api/pipelines",
        json={"name": "Bulk Tag C", "steps": [{"tier": "automation", "name": "fetch_raw_metrics"}]},
    )

    res = client.post(
        "/api/pipelines/bulk-tags",
        json={"slugs": ["bulk_tag_c", "totally_made_up_slug"], "tag": "whatever"},
    )
    assert res.status_code == 200
    assert res.json()["tagged"] == ["bulk_tag_c"]


def test_bulk_tag_pipelines_rejects_a_blank_tag():
    res = client.post("/api/pipelines/bulk-tags", json={"slugs": [], "tag": "   "})
    assert res.status_code == 400


def test_bulk_untag_pipelines_removes_the_tag_but_keeps_other_tags():
    client.post(
        "/api/pipelines",
        json={"name": "Untag A", "steps": [{"tier": "automation", "name": "fetch_raw_metrics"}]},
    )
    client.post(
        "/api/pipelines",
        json={"name": "Untag B", "steps": [{"tier": "automation", "name": "fetch_raw_metrics"}]},
    )
    client.put("/api/pipelines/untag_a/tags", json={"tags": ["keep_me", "remove_me"]})
    client.put("/api/pipelines/untag_b/tags", json={"tags": ["remove_me"]})

    res = client.post("/api/pipelines/bulk-untag", json={"slugs": ["untag_a", "untag_b"], "tag": "remove_me"})
    assert res.status_code == 200
    body = res.json()
    assert body["tag"] == "remove_me"
    assert set(body["untagged"]) == {"untag_a", "untag_b"}

    updated = {p["slug"]: p for p in client.get("/api/pipelines").json()}
    assert updated["untag_a"]["tags"] == ["keep_me"]
    assert updated["untag_b"]["tags"] == []


def test_bulk_untag_pipelines_is_a_no_op_for_a_pipeline_that_never_had_the_tag():
    client.post(
        "/api/pipelines",
        json={"name": "Untag C", "steps": [{"tier": "automation", "name": "fetch_raw_metrics"}]},
    )
    client.put("/api/pipelines/untag_c/tags", json={"tags": ["something_else"]})

    res = client.post("/api/pipelines/bulk-untag", json={"slugs": ["untag_c"], "tag": "never_had_this"})
    assert res.status_code == 200
    assert res.json()["untagged"] == ["untag_c"]

    updated = {p["slug"]: p for p in client.get("/api/pipelines").json()}
    assert updated["untag_c"]["tags"] == ["something_else"]


def test_bulk_untag_pipelines_skips_unknown_slugs_without_failing():
    client.post(
        "/api/pipelines",
        json={"name": "Untag D", "steps": [{"tier": "automation", "name": "fetch_raw_metrics"}]},
    )

    res = client.post(
        "/api/pipelines/bulk-untag",
        json={"slugs": ["untag_d", "totally_made_up_slug"], "tag": "whatever"},
    )
    assert res.status_code == 200
    assert res.json()["untagged"] == ["untag_d"]


def test_bulk_untag_pipelines_rejects_a_blank_tag():
    res = client.post("/api/pipelines/bulk-untag", json={"slugs": [], "tag": "   "})
    assert res.status_code == 400


def test_rename_pipeline_tag_relabels_it_on_every_pipeline_that_carries_it():
    client.post(
        "/api/pipelines",
        json={"name": "Rename Tag A", "steps": [{"tier": "automation", "name": "fetch_raw_metrics"}]},
    )
    client.post(
        "/api/pipelines",
        json={"name": "Rename Tag B", "steps": [{"tier": "automation", "name": "fetch_raw_metrics"}]},
    )
    client.put("/api/pipelines/rename_tag_a/tags", json={"tags": ["typo_tag", "keep_me"]})
    client.put("/api/pipelines/rename_tag_b/tags", json={"tags": ["unrelated"]})

    res = client.post("/api/pipelines/rename-tag", json={"old_tag": "typo_tag", "new_tag": "fixed_tag"})
    assert res.status_code == 200
    body = res.json()
    assert body == {"old_tag": "typo_tag", "new_tag": "fixed_tag", "renamed": ["rename_tag_a"]}

    updated = {p["slug"]: p for p in client.get("/api/pipelines").json()}
    assert set(updated["rename_tag_a"]["tags"]) == {"fixed_tag", "keep_me"}
    assert updated["rename_tag_b"]["tags"] == ["unrelated"]


def test_rename_pipeline_tag_merges_into_an_existing_new_tag_without_duplicating():
    client.post(
        "/api/pipelines",
        json={"name": "Rename Tag Merge", "steps": [{"tier": "automation", "name": "fetch_raw_metrics"}]},
    )
    client.put("/api/pipelines/rename_tag_merge/tags", json={"tags": ["old_name", "already_there"]})

    res = client.post("/api/pipelines/rename-tag", json={"old_tag": "old_name", "new_tag": "already_there"})
    assert res.status_code == 200

    updated = {p["slug"]: p for p in client.get("/api/pipelines").json()}
    assert updated["rename_tag_merge"]["tags"] == ["already_there"]


def test_rename_pipeline_tag_rejects_blank_or_identical_tags():
    res = client.post("/api/pipelines/rename-tag", json={"old_tag": "", "new_tag": "x"})
    assert res.status_code == 400

    res2 = client.post("/api/pipelines/rename-tag", json={"old_tag": "same", "new_tag": "same"})
    assert res2.status_code == 400


def test_pipeline_tags_summary_counts_tagged_pipelines():
    client.post(
        "/api/pipelines",
        json={"name": "Tagdir Pipeline A", "steps": [{"tier": "automation", "name": "fetch_raw_metrics"}]},
    )
    client.post(
        "/api/pipelines",
        json={"name": "Tagdir Pipeline B", "steps": [{"tier": "automation", "name": "fetch_raw_metrics"}]},
    )
    client.put("/api/pipelines/tagdir_pipeline_a/tags", json={"tags": ["reviewed", "q1"]})
    client.put("/api/pipelines/tagdir_pipeline_b/tags", json={"tags": ["reviewed"]})

    summary = client.get("/api/pipelines/tags-summary").json()
    by_tag = {entry["tag"]: entry["count"] for entry in summary}
    assert by_tag["reviewed"] == 2
    assert by_tag["q1"] == 1


def test_pipeline_tags_summary_csv_export_has_a_header_and_matches_the_json_summary():
    client.post(
        "/api/pipelines",
        json={"name": "Tagdir Csv Pipeline", "steps": [{"tier": "automation", "name": "fetch_raw_metrics"}]},
    )
    client.put("/api/pipelines/tagdir_csv_pipeline/tags", json={"tags": ["csv_export_marker"]})

    res = client.get("/api/pipelines/tags-summary.csv")
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("text/csv")
    assert "attachment; filename=pipeline_tags.csv" in res.headers["content-disposition"]

    lines = res.text.strip().splitlines()
    assert lines[0] == "tag,count"
    assert any(line.startswith("csv_export_marker,") for line in lines[1:])


def test_pipeline_tags_summary_json_export_matches_the_plain_json_summary():
    client.post(
        "/api/pipelines",
        json={"name": "Tagdir Json Pipeline", "steps": [{"tier": "automation", "name": "fetch_raw_metrics"}]},
    )
    client.put("/api/pipelines/tagdir_json_pipeline/tags", json={"tags": ["json_export_marker"]})

    res = client.get("/api/pipelines/tags-summary.json")
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("application/json")
    assert "attachment; filename=pipeline_tags.json" in res.headers["content-disposition"]

    rows = res.json()
    assert rows == client.get("/api/pipelines/tags-summary").json()
    assert any(entry["tag"] == "json_export_marker" for entry in rows)


def test_pipeline_tags_summary_excludes_tags_on_deleted_pipelines():
    client.post(
        "/api/pipelines",
        json={"name": "Tagdir Deleted Pipeline", "steps": [{"tier": "automation", "name": "fetch_raw_metrics"}]},
    )
    client.put("/api/pipelines/tagdir_deleted_pipeline/tags", json={"tags": ["soon_deleted"]})
    client.delete("/api/pipelines/tagdir_deleted_pipeline")

    summary = client.get("/api/pipelines/tags-summary").json()
    assert not any(entry["tag"] == "soon_deleted" for entry in summary)


def test_run_history_search_finds_step_outputs_and_errors():
    # Produce a step whose output contains a distinctive marker value.
    res = client.post(
        "/api/modules/automation/fetch_raw_metrics/run",
        json={"inputs": {"signups": 987654, "churn": 1, "revenue": 10}, "force_refresh": True},
    )
    with client.stream("GET", f"/api/stream/{res.json()['stream_id']}") as response:
        for line in response.iter_lines():
            if line.startswith("data: ") and '"run_completed"' in line:
                break

    hits = client.get("/api/runs/search", params={"q": "987654"}).json()
    assert hits["query"] == "987654"
    assert len(hits["results"]) >= 1
    top = hits["results"][0]
    assert top["step"] == "fetch_raw_metrics"
    assert top["matched_in"] == "output"
    assert "987654" in top["snippet"]
    assert top["success"] is True
    assert isinstance(top["run_id"], int)


def test_run_history_search_matches_error_text_of_failed_steps():
    res = client.post(
        "/api/modules/automation/http_request/run",
        json={"inputs": {"url": "http://127.0.0.1:9/search-marker-path", "timeout": 2}},
    )
    with client.stream("GET", f"/api/stream/{res.json()['stream_id']}") as response:
        for line in response.iter_lines():
            if line.startswith("data: ") and ('"run_completed"' in line or '"run_failed"' in line):
                break

    hits = client.get("/api/runs/search", params={"q": "search-marker-path"}).json()
    assert len(hits["results"]) >= 1
    top = hits["results"][0]
    assert top["step"] == "http_request"
    assert top["matched_in"] == "error"
    assert top["success"] is False

    # Leave no failure streak behind for other tests.
    from webapp.main import store as main_store
    main_store.reset_breaker("automation", "http_request")


def test_run_history_search_with_blank_query_returns_nothing():
    res = client.get("/api/runs/search", params={"q": "   "})
    assert res.json()["results"] == []


def test_environment_view_lists_declared_vars_and_settings(monkeypatch):
    monkeypatch.setenv("LOG_LEVEL", "DEBUG")
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-test-1234567890")

    body = client.get("/api/environment").json()

    env = {entry["name"]: entry for entry in body["environment"]}
    # Everything declared in .env.example shows up.
    assert {"ANTHROPIC_API_KEY", "DATABASE_URL", "LOG_LEVEL"} <= set(env)

    # Non-secret set var shows its actual value.
    assert env["LOG_LEVEL"]["set"] is True
    assert env["LOG_LEVEL"]["secret"] is False
    assert env["LOG_LEVEL"]["preview"] == "DEBUG"

    # Secret-shaped set var reveals only its length, never the value.
    key_entry = env["ANTHROPIC_API_KEY"]
    assert key_entry["set"] is True
    assert key_entry["secret"] is True
    assert "sk-test" not in (key_entry["preview"] or "")
    assert str(len("sk-test-1234567890")) in key_entry["preview"]

    # Unset var is reported as missing with no preview.
    assert env["DATABASE_URL"]["set"] is False
    assert env["DATABASE_URL"]["preview"] is None

    # Runtime settings reflect live config values.
    settings = {s["name"]: s["value"] for s in body["settings"]}
    assert settings["Circuit breaker threshold"] == "3"
    assert settings["Run rate limit"] == "30 requests / 10 s"
    assert settings["Upload / ingest size limit"] == "20 MB"
    assert "orchestrator.db" in settings["State store"]


def test_run_detail_endpoint_includes_recorded_inputs():
    res = client.post(
        "/api/modules/automation/fetch_raw_metrics/run",
        json={"inputs": {"signups": 321, "churn": 8, "revenue": 99}, "force_refresh": True},
    )
    with client.stream("GET", f"/api/stream/{res.json()['stream_id']}") as response:
        for line in response.iter_lines():
            if line.startswith("data: ") and '"run_completed"' in line:
                break

    run_id = client.get("/api/runs?limit=1").json()[0]["id"]
    detail = client.get(f"/api/runs/{run_id}").json()
    assert detail["steps"][0]["inputs"]["signups"] == 321
    assert detail["steps"][0]["inputs"]["churn"] == 8


def test_run_detail_endpoint_includes_an_empty_blackboard_when_nothing_was_posted():
    res = client.post(
        "/api/modules/automation/fetch_raw_metrics/run",
        json={"inputs": {"signups": 5, "churn": 1, "revenue": 1}, "force_refresh": True},
    )
    with client.stream("GET", f"/api/stream/{res.json()['stream_id']}") as response:
        for line in response.iter_lines():
            if line.startswith("data: ") and '"run_completed"' in line:
                break

    run_id = client.get("/api/runs?limit=1").json()[0]["id"]
    detail = client.get(f"/api/runs/{run_id}").json()
    assert detail["blackboard"] == []


def test_rerun_replays_a_single_module_run_with_its_recorded_inputs():
    res = client.post(
        "/api/modules/automation/fetch_raw_metrics/run",
        json={"inputs": {"signups": 654, "churn": 12, "revenue": 8}, "force_refresh": True},
    )
    with client.stream("GET", f"/api/stream/{res.json()['stream_id']}") as response:
        for line in response.iter_lines():
            if line.startswith("data: ") and '"run_completed"' in line:
                break

    run_id = client.get("/api/runs?limit=1").json()[0]["id"]
    rerun_res = client.post(f"/api/runs/{run_id}/rerun")
    assert rerun_res.status_code == 200
    assert rerun_res.json()["replayed_steps"] == 1

    events = []
    with client.stream("GET", f"/api/stream/{rerun_res.json()['stream_id']}") as response:
        for line in response.iter_lines():
            if line.startswith("data: "):
                events.append(json.loads(line[len("data: "):]))
                if events[-1]["kind"] in ("run_completed", "run_failed"):
                    break

    assert events[-1]["kind"] == "run_completed"
    assert events[-1]["context"]["signups"] == 654
    assert events[-1]["context"]["churn"] == 12


def test_rerun_replays_a_multi_step_pipeline_run_faithfully():
    res = client.post("/api/pipeline/run", json={"inputs": {}})
    with client.stream("GET", f"/api/stream/{res.json()['stream_id']}") as response:
        for line in response.iter_lines():
            if line.startswith("data: ") and '"run_completed"' in line:
                break

    run_id = client.get("/api/runs?limit=1").json()[0]["id"]
    original_steps = client.get(f"/api/runs/{run_id}").json()["steps"]
    assert len(original_steps) == 3

    rerun_res = client.post(f"/api/runs/{run_id}/rerun")
    assert rerun_res.status_code == 200
    assert rerun_res.json()["replayed_steps"] == 3

    events = []
    with client.stream("GET", f"/api/stream/{rerun_res.json()['stream_id']}") as response:
        for line in response.iter_lines():
            if line.startswith("data: "):
                events.append(json.loads(line[len("data: "):]))
                if events[-1]["kind"] in ("run_completed", "run_failed"):
                    break

    assert events[-1]["kind"] == "run_completed"
    assert "agent_decision" in events[-1]["context"]


def test_rerun_of_a_nonexistent_run_is_a_404():
    res = client.post("/api/runs/999999999/rerun")
    assert res.status_code == 404


def test_state_store_migrates_a_pre_existing_db_missing_the_inputs_column(tmp_path):
    import sqlite3

    from engine.state_store import StateStore

    db_path = str(tmp_path / "legacy.db")
    # Simulate a database created before the `inputs` column existed.
    legacy_conn = sqlite3.connect(db_path)
    legacy_conn.executescript(
        """
        CREATE TABLE runs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            started_at TEXT NOT NULL,
            finished_at TEXT,
            status TEXT NOT NULL
        );
        CREATE TABLE steps (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            run_id INTEGER NOT NULL REFERENCES runs(id),
            name TEXT NOT NULL,
            tier TEXT NOT NULL,
            success INTEGER NOT NULL,
            output TEXT NOT NULL,
            error TEXT,
            started_at TEXT NOT NULL,
            finished_at TEXT NOT NULL
        );
        """
    )
    legacy_conn.execute(
        "INSERT INTO runs (started_at, finished_at, status) VALUES ('2020-01-01T00:00:00', '2020-01-01T00:00:01', 'completed')"
    )
    legacy_conn.execute(
        "INSERT INTO steps (run_id, name, tier, success, output, started_at, finished_at) "
        "VALUES (1, 'legacy_step', 'automation', 1, '{}', '2020-01-01T00:00:00', '2020-01-01T00:00:01')"
    )
    legacy_conn.commit()
    legacy_conn.close()

    store = StateStore(db_path)  # must not raise, and must add the missing column
    steps = store.steps_for_run(1)
    assert steps[0]["name"] == "legacy_step"
    assert steps[0]["inputs"] == "{}"  # migrated default, never NULL
    store.close()

    # Reopening again (column already present) must also be a no-op, not an error.
    store2 = StateStore(db_path)
    store2.close()
